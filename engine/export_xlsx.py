"""Export Excel : uniquement une mise en forme du resultat.

Aucune formule Excel, aucune regle metier ici. Le classeur produit
est un rapport mort : ce que Python a calcule, rien de plus. C'est
precisement ce qui manquait au classeur Google Sheets.
"""

from __future__ import annotations

from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine.rules import LIBELLES_REGIME, LIBELLES_REPAS

_ENTETE = Font(bold=True, color="FFFFFF")
_FOND = PatternFill("solid", fgColor="2F5597")
_GRAS = Font(bold=True)


def _feuille(classeur: Workbook, titre: str, colonnes: list[str]):
    feuille = classeur.create_sheet(titre)
    feuille.append(colonnes)
    for cellule in feuille[1]:
        cellule.font = _ENTETE
        cellule.fill = _FOND
        cellule.alignment = Alignment(horizontal="center")
    feuille.freeze_panes = "A2"
    return feuille


def _formater_dates(feuille) -> None:
    """Sans cela, openpyxl ecrit les dates comme des datetime bruts."""
    for ligne in feuille.iter_rows(min_row=2, min_col=1, max_col=1):
        for cellule in ligne:
            if hasattr(cellule.value, "year"):
                cellule.number_format = "DD/MM/YYYY"


def _ajuster(feuille) -> None:
    for colonne in feuille.columns:
        largeur = max((len(str(c.value)) for c in colonne if c.value is not None), default=8)
        feuille.column_dimensions[get_column_letter(colonne[0].column)].width = min(largeur + 3, 40)


def exporter(chemin, personnes: dict, prestations, facturation, config: dict) -> None:
    devise = config.get("meta", {}).get("devise", "EUR")
    classeur = Workbook()
    classeur.remove(classeur.active)

    # --- 1. Ce que l'hotel a besoin de savoir : combien de quoi, quel jour ---
    regimes = list(LIBELLES_REGIME)
    feuille = _feuille(classeur, "Synthese hotel", ["Date"] + [LIBELLES_REGIME[r] for r in regimes])
    compte = Counter((n.jour, n.regime) for n in prestations.nuitees)
    for jour in sorted({n.jour for n in prestations.nuitees}):
        feuille.append([jour] + [compte[(jour, r)] for r in regimes])
    ligne_total = ["TOTAL"] + [
        sum(n for (_, reg), n in compte.items() if reg == r) for r in regimes
    ]
    feuille.append(ligne_total)
    for cellule in feuille[feuille.max_row]:
        cellule.font = _GRAS
    _formater_dates(feuille)
    _ajuster(feuille)

    # --- 2. Repas servis hors pension ---
    feuille = _feuille(classeur, "Repas hors pension", ["Date", "Repas", "Nombre"])
    compte_repas = Counter((r.jour, r.repas) for r in prestations.repas_hors_pension)
    for (jour, repas), nombre in sorted(compte_repas.items()):
        feuille.append([jour, LIBELLES_REPAS[repas], nombre])
    _formater_dates(feuille)
    _ajuster(feuille)

    # --- 3. Detail par personne ---
    feuille = _feuille(
        classeur,
        "Detail personne",
        ["Date", "Famille", "Prenom", "Age", "Prestation", "Detail", f"Prix ({devise})"],
    )
    for ligne in facturation.lignes:
        personne = personnes[ligne.personne_id]
        feuille.append(
            [
                ligne.jour,
                personne.famille,
                personne.prenom,
                personne.categorie_age,
                ligne.libelle,
                ligne.detail,
                ligne.prix,
            ]
        )
    _formater_dates(feuille)
    _ajuster(feuille)

    # --- 4. Total par famille ---
    feuille = _feuille(classeur, "Par famille", ["Famille", "Personnes", f"Total ({devise})"])
    totaux = defaultdict(float)
    membres = defaultdict(set)
    for ligne in facturation.lignes:
        personne = personnes[ligne.personne_id]
        totaux[personne.famille] += ligne.prix
        membres[personne.famille].add(personne.prenom)
    for famille in sorted(totaux):
        feuille.append([famille, len(membres[famille]), round(totaux[famille], 2)])
    feuille.append(["TOTAL", sum(len(m) for m in membres.values()), facturation.total])
    for cellule in feuille[feuille.max_row]:
        cellule.font = _GRAS
    _formater_dates(feuille)
    _ajuster(feuille)

    # --- 5. Ce qu'il reste a demander a l'hotel ---
    feuille = _feuille(
        classeur,
        "Tarifs manquants",
        ["Prix a zero ou absent - a verifier dans l'onglet Tarifs"],
    )
    for cle in sorted(facturation.tarifs_manquants):
        feuille.append([cle])
    if not facturation.tarifs_manquants:
        feuille.append(["Aucun : tous les tarifs utilises sont renseignes."])
    _formater_dates(feuille)
    _ajuster(feuille)

    # --- 6. Saisies incoherentes ---
    feuille = _feuille(classeur, "Anomalies", ["Date", "Famille", "Prenom", "Probleme"])
    for anomalie in prestations.anomalies:
        personne = personnes[anomalie.personne_id]
        feuille.append([anomalie.jour, personne.famille, personne.prenom, anomalie.message])

    # Un gite se loue entier : sans place attribuee, on ne sait pas avec
    # combien la note se partage, et la nuit reste a zero. Ce n'est pas une
    # erreur de saisie -- c'est un plan de couchage a finir -- mais ca se
    # lit au meme endroit, sinon ca ne se lit nulle part.
    for personne_id, jour in facturation.sans_place:
        personne = personnes[personne_id]
        feuille.append([
            jour, personne.famille, personne.prenom,
            "En gite sans place attribuee : nuit non facturee",
        ])

    if not prestations.anomalies and not facturation.sans_place:
        feuille.append(["", "", "", "Aucune anomalie detectee."])
    _formater_dates(feuille)
    _ajuster(feuille)

    classeur.save(chemin)
