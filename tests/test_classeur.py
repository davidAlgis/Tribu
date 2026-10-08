"""Le classeur que la page des comptes fait telecharger s'ouvre-t-il ?

`classeur.js` ecrit un .xlsx a la main -- une archive zip de quelques
fichiers XML -- pour ne pas faire dependre une page de famille d'une
bibliotheque chargee depuis un autre serveur. Ecrire un format a la main,
c'est pouvoir se tromper d'un octet sans que rien ne le dise : le fichier
se telecharge, et Excel refuse de l'ouvrir chez quelqu'un d'autre.

Ce test le fait ecrire par Node et le relit avec openpyxl, la bibliotheque
qu'utilise deja l'export Python. Sans `node`, il se saute.
"""

import io
import json
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

RACINE = Path(__file__).resolve().parent.parent

FEUILLES = [
    {
        "nom": "Ce que chacun paie",
        "lignes": [
            ["Nom Prénom", "Hébergement", "Suppl. / réd.", "Total"],
            ["Alice (A)", 81, -20.25, 60.75],
            ['Bruno <&> "B"', 61, 15, 76],
            ["TOTAL", 142, None, 136.75],
        ],
    },
    # Un nom trop long, et des caracteres qu'Excel refuse.
    {"nom": "Les couverts, repas par repas: [tous] / les jours", "lignes": [["Repas"], ["Dîner 10"]]},
    # Deux feuilles du meme nom.
    {"nom": "Ce que chacun paie", "lignes": []},
]


def ecrire(script: str) -> bytes:
    sortie = subprocess.run(
        ["node", "-e", script], capture_output=True, check=True
    )
    return sortie.stdout


def charger_module() -> str:
    return (
        "const fs = require('fs'); const fenetre = {};"
        f"new Function('window', fs.readFileSync({str(RACINE / 'classeur.js')!r}, 'utf8'))(fenetre);"
    )


@pytest.fixture(scope="module")
def classeur():
    if not shutil.which("node"):
        pytest.skip("node absent : le classeur ne s'ecrit pas")
    octets = ecrire(
        charger_module()
        + f"process.stdout.write(Buffer.from(fenetre.CLASSEUR.xlsx({json.dumps(FEUILLES)})));"
    )
    return openpyxl.load_workbook(io.BytesIO(octets))


def test_chaque_tableau_a_sa_feuille(classeur):
    assert len(classeur.sheetnames) == 3
    assert classeur.sheetnames[0] == "Ce que chacun paie"
    # Trente et un caracteres au plus, sans ce qu'Excel refuse.
    long = classeur.sheetnames[1]
    assert len(long) <= 31 and not set(long) & set("[]:*?/\\")
    # Deux feuilles ne portent pas le meme nom.
    assert classeur.sheetnames[2] != classeur.sheetnames[0]


def test_un_montant_reste_un_nombre(classeur):
    """« 60,75 € » se somme dans Excel ; son texte, non."""
    feuille = classeur["Ce que chacun paie"]
    lignes = [[c.value for c in r] for r in feuille.iter_rows()]
    assert lignes[0] == ["Nom Prénom", "Hébergement", "Suppl. / réd.", "Total"]
    assert lignes[1] == ["Alice (A)", 81, -20.25, 60.75]
    assert lignes[2][0] == 'Bruno <&> "B"'
    # Une case vide reste vide, pas « None » ni zero.
    assert lignes[3] == ["TOTAL", 142, None, 136.75]
    assert feuille["A1"].font.bold


def test_le_csv_s_ouvre_dans_un_excel_francais():
    """Point-virgule entre les cases, virgule decimale, BOM UTF-8 : sans
    eux, Excel en francais met tout dans la premiere colonne, ou casse les
    accents."""
    if not shutil.which("node"):
        pytest.skip("node absent")
    texte = ecrire(
        charger_module()
        + "process.stdout.write(fenetre.CLASSEUR.csv("
        + json.dumps([["Nom", "Total"], ["Un texte; « entre guillemets »", 12.5], ["Vide", None]])
        + "));"
    ).decode("utf-8")
    assert texte.startswith("﻿")
    assert texte.split("\r\n")[:3] == [
        "﻿Nom;Total",
        '"Un texte; « entre guillemets »";12,5',
        "Vide;",
    ]
